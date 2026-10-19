'''
Unit tests for `evaluator plot`
'''

# ====================
# Import external dependencies
# ====================
import pandas as pd, plotly.graph_objects as go, pytest, tomllib
from pathlib import Path
from unittest.mock import MagicMock, patch

# ====================
# Import evaluator plot functions
# ====================
from evaluator.commands.plot.plot import run_plot
from evaluator.commands.plot.utils import data as datautil
from evaluator.commands.plot.utils import export as exportutil
from evaluator.commands.plot.utils.input import (
    PlotRun,
    available_sections,
    parse_agreement,
    resolve_plot_inputs,
)

# ====================
# Define helpers
# ====================
def _write(path: Path, text: str) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding='utf-8')
    return path

def _sheet(path: Path, rows: list[dict]) -> Path:
    header = ['sample_id', 'path', 'group', 'replicate']
    lines = ['\t'.join(header)]
    for row in rows:
        lines.append('\t'.join(str(row.get(col, '')) for col in header))
    return _write(path, '\n'.join(lines) + '\n')

def _make_analyse_csv(tmp_path):
    path = tmp_path / 'analyse.csv'
    path.write_text(
        'tomogram,label,equiv_diameter_nm,major_axis_diameter,minor_axis_diameter,'
        'aspect_ratio,eccentricity,membrane_volume,lumen_volume,surface_area,'
        'is_enclosed,closure_fill_ratio\n'
        'tomo1.mrc,1,80.0,90.0,70.0,1.28,0.5,50000,30000,20000,True,0.4\n'
    )
    return path

@pytest.fixture
def fake_sections():
    names = ['distributions', 'qc', 'scatter', 'correlations', 'concordance', 'reliability', 'agreement']
    mocks = {n: MagicMock() for n in names}
    with patch.dict('evaluator.commands.plot.sections.SECTIONS', mocks, clear=True), patch('evaluator.commands.plot.plot.exportutil.check_kaleido'), patch('evaluator.commands.plot.plot.sectionutil.load'):
        yield mocks

# ====================
# Define tests for plot/utils/input.py
# ====================
class TestPlotUtilsInput:
    # == Define tests for resolve_plot_inputs with single file (no sample sheet) ==
    def test_analyse_only(self, tmp_path):
        analyse = _write(tmp_path / 'analyse.csv', 'tomogram,label\n')
        runs, multi_run, sheet_path = resolve_plot_inputs(analyse, None)
        assert multi_run is False
        assert sheet_path is None
        assert runs == [PlotRun(sample_id='sample', analyse_path=analyse, model_path=None)]

    def test_model_only(self, tmp_path):
        model = _write(tmp_path / 'model.json', '{}')
        runs, multi_run, sheet_path = resolve_plot_inputs(None, model)
        assert multi_run is False
        assert runs == [PlotRun(sample_id='sample', analyse_path=None, model_path=model)]

    def test_both_present(self, tmp_path):
        analyse = _write(tmp_path / 'analyse.csv', 'tomogram,label\n')
        model = _write(tmp_path / 'model.json', '{}')
        runs, multi_run, sheet_path = resolve_plot_inputs(analyse, model)
        assert multi_run is False
        assert runs == [PlotRun(sample_id='sample', analyse_path=analyse, model_path=model)]

    def test_csv_suffix_never_treated_as_sheet(self, tmp_path):
        '''A .csv with a sample_id-like header must not be sniffed as a sheet (only .tsv/.txt are)'''
        analyse = _write(tmp_path / 'analyse.csv', 'sample_id\tpath\n')
        runs, multi_run, sheet_path = resolve_plot_inputs(analyse, None)
        assert multi_run is False

    # == Define tests for resolve_plot_inputs with a sample sheet (from analyse or model) ==
    def test_analyse_sheet_only(self, tmp_path):
        a1 = _write(tmp_path / 'a1.csv', 'tomogram,label\n')
        a2 = _write(tmp_path / 'a2.csv', 'tomogram,label\n')
        sheet = _sheet(tmp_path / 'sheet.tsv', [
            {'sample_id': 's1', 'path': a1, 'group': 'control'},
            {'sample_id': 's2', 'path': a2, 'group': 'treated'},
        ])
        runs, multi_run, sheet_path = resolve_plot_inputs(sheet, None)
        assert multi_run is True
        assert sheet_path == sheet
        assert {r.sample_id for r in runs} == {'s1', 's2'}
        assert {r.group for r in runs} == {'control', 'treated'}

    def test_model_sheet_only(self, tmp_path):
        m1 = _write(tmp_path / 'm1.json', '{}')
        sheet = _sheet(tmp_path / 'sheet.tsv', [{'sample_id': 's1', 'path': m1, 'group': 'g'}])
        runs, multi_run, sheet_path = resolve_plot_inputs(None, sheet)
        assert multi_run is True
        assert runs[0].model_path == m1
        assert runs[0].analyse_path is None

    def test_both_sheets_joined_on_sample_id(self, tmp_path):
        a1 = _write(tmp_path / 'a1.csv', 'tomogram,label\n')
        m1 = _write(tmp_path / 'm1.json', '{}')
        analyse_sheet = _sheet(tmp_path / 'analyse_sheet.tsv', [{'sample_id': 's1', 'path': a1, 'group': 'g'}])
        model_sheet = _sheet(tmp_path / 'model_sheet.tsv', [{'sample_id': 's1', 'path': m1, 'group': 'g'}])
        runs, multi_run, sheet_path = resolve_plot_inputs(analyse_sheet, model_sheet)
        assert multi_run is True
        # analyse sheet is preferred as the returned sheet_path when both sides are sheets
        assert sheet_path == analyse_sheet
        assert len(runs) == 1
        assert runs[0].analyse_path == a1
        assert runs[0].model_path == m1

    def test_replicate_parsed_as_int(self, tmp_path):
        a1 = _write(tmp_path / 'a1.csv', 'tomogram,label\n')
        sheet = _sheet(tmp_path / 'sheet.tsv', [{'sample_id': 's1', 'path': a1, 'group': 'g', 'replicate': 2}])
        runs, _, _ = resolve_plot_inputs(sheet, None)
        assert runs[0].replicate == 2

    def test_missing_sample_id_column_raises(self, tmp_path):
        bad = _write(tmp_path / 'sheet.tsv', 'path\tgroup\n/x.csv\tg\n')
        with pytest.raises(ValueError):
            resolve_plot_inputs(bad, None)

    # == Define tests for available_sections ==
    def test_analyse_only_gives_core_sections(self):
        runs = [PlotRun(sample_id='s', analyse_path=Path('a.csv'), model_path=None)]
        assert available_sections(runs, False) == ['distributions', 'qc', 'scatter', 'correlations']

    def test_model_only_gives_no_sections(self):
        '''No section currently runs on model output alone'''
        runs = [PlotRun(sample_id='s', analyse_path=None, model_path=Path('m.json'))]
        assert available_sections(runs, False) == []

    def test_analyse_and_model_adds_concordance(self):
        runs = [PlotRun(sample_id='s', analyse_path=Path('a.csv'), model_path=Path('m.json'))]
        sections = available_sections(runs, False)
        assert 'concordance' in sections
        assert 'reliability' in sections
        assert 'agreement' in sections

    def test_single_group_excludes_compare(self):
        runs = [
            PlotRun(sample_id='s1', analyse_path=Path('a1.csv'), model_path=None, group='g'),
            PlotRun(sample_id='s2', analyse_path=Path('a2.csv'), model_path=None, group='g'),
        ]
        assert 'compare' not in available_sections(runs, True)

    def test_multiple_groups_adds_compare(self):
        runs = [
            PlotRun(sample_id='s1', analyse_path=Path('a1.csv'), model_path=None, group='control'),
            PlotRun(sample_id='s2', analyse_path=Path('a2.csv'), model_path=None, group='treated'),
        ]
        assert 'compare' in available_sections(runs, True)

    def test_single_run_mode_never_adds_compare(self):
        '''compare requires multi_run even if (somehow) groups differ'''
        runs = [PlotRun(sample_id='s', analyse_path=Path('a.csv'), model_path=None, group='g')]
        assert 'compare' not in available_sections(runs, False)

    def test_parse_agreement_repeatable_and_space_separated(self):
        assert parse_agreement(['a,b', 'c,d e,f']) == [('a', 'b'), ('c', 'd'), ('e', 'f')]

    @pytest.mark.parametrize('bad', ['a', 'a,b,c', ',b', 'a,'])
    def test_parse_agreement_rejects_malformed_pairs(self, bad):
        with pytest.raises(ValueError):
            parse_agreement([bad])

# ====================
# Define tests for plot/plot.py
# ====================
class TestPlot:
    # == Define tests for run_plot section gating ==
    def test_requested_but_unrunnable_section_is_skipped(self, tmp_path, fake_sections):
        analyse = _make_analyse_csv(tmp_path)
        run_plot(analyse, None, tmp_path, sections=['concordance'], all_sections=False, overwrite=False)
        fake_sections['concordance'].assert_not_called()

    def test_all_sections_runs_every_runnable_section(self, tmp_path, fake_sections):
        analyse = _make_analyse_csv(tmp_path)
        run_plot(analyse, None, tmp_path, sections=None, all_sections=True, overwrite=False)
        called = {n for n, m in fake_sections.items() if m.called}
        assert called == {'distributions', 'qc', 'scatter', 'correlations'}

    def test_include_flagged_passed_to_loader(self, tmp_path, fake_sections):
        analyse = _make_analyse_csv(tmp_path)
        with patch('evaluator.commands.plot.plot.sectionutil.load') as mock_load:
            run_plot(analyse, None, tmp_path, sections=['qc'], all_sections=False, overwrite=False, include_flagged=True)
        assert mock_load.call_args.kwargs['include_flagged'] is True

    def test_section_failure_is_not_fatal(self, tmp_path, fake_sections):
        analyse = _make_analyse_csv(tmp_path)
        fake_sections['distributions'].side_effect = ValueError('boom')
        run_plot(analyse, None, tmp_path, sections=['distributions', 'qc'], all_sections=False, overwrite=False)  # must not raise
        fake_sections['qc'].assert_called_once()

    # == Define tests for multi-sample output ==
    def test_multi_sample_writes_subdirectory_per_sample(self, tmp_path, fake_sections):
        a1, a2 = _make_analyse_csv(tmp_path), _write(tmp_path / 'a2.csv', 'x')
        sheet = _sheet(tmp_path / 'sheet.tsv', [{'sample_id': 's1', 'path': a1}, {'sample_id': 's2', 'path': a2}])
        run_plot(sheet, None, tmp_path, sections=['qc'], all_sections=False, overwrite=False)
        assert fake_sections['qc'].call_count == 2
        qc_dir = tmp_path / 'evaluator' / 'plot' / 'qc'
        assert (qc_dir / 's1').is_dir() and (qc_dir / 's2').is_dir()

    # == Define tests for --overwrite option ==
    def test_existing_section_dir_skipped_without_overwrite(self, tmp_path, fake_sections):
        analyse = _make_analyse_csv(tmp_path)
        run_plot(analyse, None, tmp_path, sections=['qc'], all_sections=False, overwrite=False)
        fake_sections['qc'].reset_mock()
        run_plot(analyse, None, tmp_path, sections=['qc'], all_sections=False, overwrite=False)
        fake_sections['qc'].assert_not_called()

    def test_overwrite_reruns_existing_section(self, tmp_path, fake_sections):
        analyse = _make_analyse_csv(tmp_path)
        run_plot(analyse, None, tmp_path, sections=['qc'], all_sections=False, overwrite=False)
        fake_sections['qc'].reset_mock()
        run_plot(analyse, None, tmp_path, sections=['qc'], all_sections=False, overwrite=True)
        fake_sections['qc'].assert_called_once()

    # == Define tests for _write_index and params ==
    def test_index_md_lists_completed_sections(self, tmp_path, fake_sections):
        analyse = _make_analyse_csv(tmp_path)
        run_plot(analyse, None, tmp_path, sections=['qc'], all_sections=False, overwrite=False)
        assert 'qc/' in (tmp_path / 'evaluator' / 'plot' / 'index.md').read_text()

    def test_params_toml_written(self, tmp_path, fake_sections):
        analyse = _make_analyse_csv(tmp_path)
        run_plot(analyse, None, tmp_path, sections=['qc'], all_sections=False, overwrite=False)
        params = tomllib.loads((tmp_path / 'evaluator' / 'plot' / 'params.toml').read_text())
        assert params['fill_ratio_flag_threshold'] == 0.05 and params['theme'] == 'Okabe-Ito'

    def test_agreement_pairs_add_the_section(self, tmp_path, fake_sections):
        analyse = _make_analyse_csv(tmp_path)
        model = _write(tmp_path / 'model.json', '{}')
        run_plot(analyse, model, tmp_path, sections=['qc'], all_sections=False, overwrite=False, agreement=[('a', 'b')])
        fake_sections['agreement'].assert_called_once()

# ====================
# Define tests for plot/utils/data.py
# ====================
class TestPlotUtilsData:
    def test_attach_fitted_diameter_masks_unreliable(self):
        analyse = pd.DataFrame({'tomogram': ['t.mrc'] * 3, 'label': [1, 2, 3]})
        model = pd.DataFrame({
            'source_file': ['/x/t.mrc', '/x/t.mrc'], 'label_id': [1, 2], 'model_diameter_nm': [100.0, 50.0],
            'reliability.is_reliable': [True, False],
        })
        out = datautil.attach_fitted_diameter(analyse, model)
        assert out['model_diameter_nm'].tolist()[0] == 100.0
        assert out['model_diameter_nm'].isna().tolist() == [False, True, True]  # unreliable fit and unmatched row both NaN

    def test_attach_fitted_diameter_keeps_row_count_with_duplicate_model_rows(self):
        analyse = pd.DataFrame({'tomogram': ['t.mrc'], 'label': [1]})
        model = pd.DataFrame({'source_file': ['/x/t.mrc', '/y/t.mrc'], 'label_id': [1, 1], 'model_diameter_nm': [10.0, 20.0]})
        assert len(datautil.attach_fitted_diameter(analyse, model)) == 1

    def test_populations_split_by_qc_and_z_clip(self):
        df = pd.DataFrame({
            'is_vesicle_like': [True, True, False, True],
            'touches_z_boundary': [False, True, False, False],
        })
        pops = datautil.populations(df)
        assert {k: len(v) for k, v in pops.items()} == {'all': 4, 'vesicle_like': 3, 'vesicle_like_no_z_clip': 2}

    def test_populations_without_qc_column_is_all_only(self):
        assert list(datautil.populations(pd.DataFrame({'x': [1, 2]}))) == ['all']

    def _csv(self, tmp_path, rows):
        header = 'tomogram,label,equiv_diameter_nm,major_axis_diameter,minor_axis_diameter,aspect_ratio,eccentricity,membrane_volume,lumen_volume,surface_area,is_enclosed,closure_fill_ratio,is_vesicle_like,qc_flags\n'
        body = ''.join(f'tomo1.mrc,{i},80,90,70,1.2,0.5,5,3,2,True,0.4,{vl},{flags}\n' for i, (vl, flags) in enumerate(rows))
        return _write(tmp_path / 'a.csv', header + body)

    def test_flagged_rows_dropped_by_default(self, tmp_path):
        path = self._csv(tmp_path, [(True, ''), (False, 'sphere_rmse'), (False, 'sphere_rmse|x')])
        assert len(datautil.load_analyse(path)) == 1

    def test_include_flagged_keeps_everything(self, tmp_path):
        path = self._csv(tmp_path, [(True, ''), (False, 'sphere_rmse')])
        assert len(datautil.load_analyse(path, include_flagged=True)) == 2

    def test_missing_vesicle_like_column_is_not_filtered(self, tmp_path):
        path = _make_analyse_csv(tmp_path)
        assert len(datautil.load_analyse(path)) == 1

    def test_missing_required_column_raises(self, tmp_path):
        path = _write(tmp_path / 'a.csv', 'tomogram,label\ntomo1.mrc,1\n')
        with pytest.raises(ValueError, match='missing required'):
            datautil.load_analyse(path)

    def test_drop_reasons_counted_per_flag(self, tmp_path, capsys):
        from evaluator.utils.settings import logger
        messages = []
        sink = logger.add(lambda m: messages.append(str(m)), level='INFO')
        try:
            path = self._csv(tmp_path, [(True, ''), (False, '"sphere_rmse,open_shell"'), (False, 'sphere_rmse')])
            datautil.load_analyse(path)
        finally:
            logger.remove(sink)
        assert any('sphere_rmse=2' in m and 'open_shell=1' in m for m in messages)

    def test_join_matches_on_basename_and_label_and_counts_unmatched(self):
        analyse = pd.DataFrame({'tomogram': ['t.mrc', 't.mrc', 't.mrc'], 'label': [1, 2, 3]})
        model = pd.DataFrame({'source_file': ['/x/t.mrc', '/x/t.mrc'], 'label_id': [1, 9], 'radius': [10.0, 5.0]})
        joined, a_only, m_only = datautil.join_analyse_model(analyse, model)
        assert list(joined['label']) == [1] and (a_only, m_only) == (2, 1)

    def test_pooling_note_states_true_n(self):
        df = pd.DataFrame({'tomogram': ['a', 'a', 'b']})
        note = datautil.pooling_note(df, n_samples=2)
        assert 'n = 3 EVs' in note and '2 tomogram(s)' in note and '2 sample(s)' in note

# ====================
# Define tests for plot/utils/export.py
# ====================
def _kaleido_ok() -> bool:
    try:
        exportutil.check_kaleido()
        return True
    except RuntimeError:
        return False

class TestPlotUtilsExport:
    def test_vectorise_replaces_scattergl(self):
        fig = go.Figure(go.Scattergl(x=[1, 2], y=[1, 2], mode='markers'))
        assert fig.data[0].type == 'scattergl'
        assert exportutil._vectorise(fig).data[0].type == 'scatter'

    def test_write_xlsx_first_sheet_active(self, tmp_path):
        from openpyxl import load_workbook
        path = tmp_path / 't.xlsx'
        exportutil.write_xlsx(path, {'main': pd.DataFrame({'a': [1, 2]}), 'notes': pd.DataFrame({'note': ['x']})})
        wb = load_workbook(path)
        assert wb.sheetnames == ['main', 'notes'] and wb.active.title == 'main'

    @pytest.mark.skipif(not _kaleido_ok(), reason='kaleido/Chrome unavailable')
    def test_write_figure_svg(self, tmp_path):
        path = tmp_path / 'f.svg'
        exportutil.write_figure(go.Figure(go.Scatter(x=[1, 2], y=[1, 2])), path)
        assert path.read_text().lstrip().startswith(('<svg', '<?xml'))
